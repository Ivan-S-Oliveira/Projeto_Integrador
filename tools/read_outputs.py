#!/usr/bin/env python3
"""Consolida resultados locais do projeto em UM DOCX autocontido.

Execute na raiz do repositório:
    python tools/read_outputs.py

Saída única:
    .local-analysis/compiled_results/RESULTADOS_COMPLETOS.docx
"""
from __future__ import annotations
import argparse, hashlib, json, os, platform, re, subprocess, sys, tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
try: import pandas as pd
except Exception: pd=None
try: from PIL import Image
except Exception: Image=None
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TEXT={'.txt','.md','.log','.rst'}; TABLE={'.csv','.tsv'}; STRUCT={'.json','.yaml','.yml'}
IMAGES={'.png','.jpg','.jpeg','.webp','.gif','.bmp','.tif','.tiff'}
MODELS={'.joblib','.pkl','.pickle','.onnx','.pt','.pth','.h5','.hdf5'}; DATA={'.parquet','.feather','.arrow'}; NOTEBOOK={'.ipynb'}
RESULT_DIRS={'outputs','output','results','resultados','reports','report','figures','figuras','tables','tabelas','metrics','metricas','predictions','predicoes','runs','models','modelos','artifacts','artefatos','logs','evaluation','avaliacao','experiments','experimentos'}
HINTS={'result','resultado','metric','metrica','score','prediction','predicao','evaluation','avaliacao','report','relatorio','manifest','manifesto','waterfall','comparison','comparacao','calibration','calibracao','confusion','confusao','importance','importancia','bootstrap','subgroup','subgrupo','error','erro','performance','benchmark','summary','resumo','feature','cluster','silhouette','auc','brier','threshold','limiar','shap','umap','pca','curva','matriz'}
EXCLUDE={'.git','.venv','venv','env','__pycache__','.pytest_cache','.mypy_cache','.ruff_cache','.tox','.nox','.ipynb_checkpoints','node_modules','.idea','.vscode','site-packages','build','dist','.local-analysis','.srag-status-local','.project-status-local'}
SECRET=re.compile(r'(?im)^\s*([A-Za-z_][\w.-]*(?:token|secret|password|passwd|pwd|api[_-]?key|private[_-]?key|credential)[\w.-]*)\s*[:=]\s*(.+)$')
HOMEWIN=re.compile(r'(?i)\b[A-Z]:[\\/]+Users[\\/]+[^\\/\s\"\'<>|]+'); HOMEPOSIX=re.compile(r'(?<![A-Za-z0-9_])/(?:home|Users)/[^/\s\"\'<>|]+')
NAVY='17365D'; BLUE='2F5597'; GREY='F3F5F9'; AMBER='FFF4E5'; RED='FDECEA'

def repo_root(start:Path)->Path:
 p=subprocess.run(['git','rev-parse','--show-toplevel'],cwd=start,text=True,capture_output=True,encoding='utf-8',errors='replace'); return Path(p.stdout.strip()).resolve() if p.returncode==0 and p.stdout.strip() else start.resolve()
def git(root:Path,*a:str)->str:
 p=subprocess.run(['git',*a],cwd=root,text=True,capture_output=True,encoding='utf-8',errors='replace'); return p.stdout.strip() if p.returncode==0 else 'indisponível'
def sha(path:Path)->str:
 h=hashlib.sha256()
 with path.open('rb') as f:
  for c in iter(lambda:f.read(1024*1024),b''):h.update(c)
 return h.hexdigest()
def size(n:int)->str:
 v=float(n)
 for u in ['B','KB','MB','GB','TB']:
  if v<1024 or u=='TB':return f'{v:.0f} {u}' if u=='B' else f'{v:.1f} {u}'
  v/=1024
def redact(t:str,root:Path)->str:
 t=SECRET.sub(lambda m:f'{m.group(1)}=[REDACTED]',t).replace(str(root),'<REPO>').replace(root.as_posix(),'<REPO>'); return HOMEPOSIX.sub('<HOME>',HOMEWIN.sub('<HOME>',t))
def kind(p:Path)->str:
 e=p.suffix.lower()
 if e in TABLE:return 'tabela'
 if e in TEXT:return 'texto'
 if e in STRUCT:return 'estruturado'
 if e in IMAGES:return 'imagem'
 if e in MODELS:return 'modelo'
 if e in DATA:return 'dados'
 if e in NOTEBOOK:return 'notebook'
 return 'outro'
def wanted(p:Path,root:Path,notebooks:bool)->bool:
 rel=p.relative_to(root); parts=[x.lower() for x in rel.parts]
 if any(x in EXCLUDE for x in parts[:-1]):return False
 e=p.suffix.lower(); supported=e in TEXT|TABLE|STRUCT|IMAGES|MODELS|DATA or (notebooks and e in NOTEBOOK)
 return supported and (any(x in RESULT_DIRS for x in parts[:-1]) or any(h in p.stem.lower() for h in HINTS))
def scan(root:Path,notebooks:bool)->list[Path]:
 out=[]
 for cur,dirs,files in os.walk(root):
  dirs[:]=sorted(d for d in dirs if d.lower() not in EXCLUDE)
  for n in sorted(files):
   p=Path(cur)/n
   try:
    if wanted(p,root,notebooks):out.append(p)
   except Exception:pass
 return sorted(out,key=lambda p:p.relative_to(root).as_posix().lower())
def shade(c:Any,color:str):
 pr=c._tc.get_or_add_tcPr();x=pr.find(qn('w:shd'))
 if x is None:x=OxmlElement('w:shd');pr.append(x)
 x.set(qn('w:fill'),color)
def margins(c:Any,v:int=70):
 pr=c._tc.get_or_add_tcPr();m=pr.first_child_found_in('w:tcMar')
 if m is None:m=OxmlElement('w:tcMar');pr.append(m)
 for side in ['top','start','bottom','end']:
  x=m.find(qn('w:'+side))
  if x is None:x=OxmlElement('w:'+side);m.append(x)
  x.set(qn('w:w'),str(v));x.set(qn('w:type'),'dxa')
def header(row:Any):
 pr=row._tr.get_or_add_trPr();x=OxmlElement('w:tblHeader');x.set(qn('w:val'),'true');pr.append(x)
def table(doc:Document,heads:list[str],rows:list[tuple],font=7):
 t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,h in enumerate(heads):
  c=t.rows[0].cells[j];shade(c,NAVY);margins(c);r=c.paragraphs[0].add_run(str(h));r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(font)
 header(t.rows[0])
 for i,row in enumerate(rows):
  cs=t.add_row().cells
  for j,v in enumerate(row):
   margins(cs[j],55);cs[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i%2:shade(cs[j],GREY)
   r=cs[j].paragraphs[0].add_run(str(v));r.font.size=Pt(font-0.5)
 return t
def callout(doc:Document,title:str,text:str,color=AMBER):
 t=doc.add_table(rows=1,cols=1);c=t.cell(0,0);shade(c,color);margins(c,130);r=c.paragraphs[0].add_run(title+'\n');r.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY);c.paragraphs[0].add_run(text);doc.add_paragraph()
def add_text(doc:Document,text:str,root:Path,limit:int):
 text=redact(text,root)
 if len(text)>limit:text=text[:int(limit*.8)]+f'\n\n[... {len(text)-limit} caracteres omitidos ...]\n\n'+text[-int(limit*.2):]
 for b in re.split(r'\n{2,}',text):
  if b.strip():
   p=doc.add_paragraph();r=p.add_run(b.strip());r.font.name='Consolas';r.font.size=Pt(8);p.paragraph_format.space_after=Pt(3)
def read(path:Path)->str:return path.read_bytes().decode('utf-8-sig',errors='replace')
def add_csv(doc:Document,p:Path,a):
 if pd is None:add_text(doc,read(p),a.root_resolved,a.max_text_chars);return
 try:df=pd.read_csv(p,sep='\t' if p.suffix.lower()=='.tsv' else ',',encoding='utf-8-sig',low_memory=False,nrows=a.max_rows_read)
 except Exception as e:callout(doc,'FALHA AO LER TABELA',repr(e),RED);add_text(doc,read(p),a.root_resolved,a.max_text_chars);return
 table(doc,['Campo','Valor'],[('Linhas analisadas',len(df)),('Colunas',len(df.columns)),('Limite de leitura',a.max_rows_read)])
 shown=df.iloc[:a.rows_docx,:a.cols_docx];rows=[]
 for _,r in shown.iterrows():rows.append(tuple('<NA>' if pd.isna(v) else str(v)[:300] for v in r))
 if len(shown.columns):table(doc,[str(x) for x in shown.columns],rows,font=6.5)
 stats=[]
 for c in df.columns[:a.summary_cols]:
  s=df[c];miss=int(s.isna().sum())
  if pd.api.types.is_numeric_dtype(s):
   d=s.dropna().describe();detail='; '.join(f'{k}={float(d[k]):.6g}' for k in ['count','mean','std','min','25%','50%','75%','max'] if k in d and not pd.isna(d[k]))
  else:detail='; '.join(f'{k}={int(v)}' for k,v in s.astype('string').fillna('<NA>').value_counts().head(a.top_values).items())
  stats.append((c,str(s.dtype),miss,detail))
 table(doc,['Coluna','Tipo','Ausentes','Resumo'],stats,font=6.5)
def flatten(x:Any,prefix='',limit=12000):
 out=[]
 def rec(v,p):
  if len(out)>=limit:return
  if isinstance(v,dict):
   for k,z in v.items():rec(z,f'{p}.{k}' if p else str(k))
  elif isinstance(v,list):
   out.append((p,f'<lista com {len(v)} itens>'))
   for i,z in enumerate(v[:50]):rec(z,f'{p}[{i}]')
  else:out.append((p,repr(v)))
 rec(x,prefix);return out
def add_struct(doc:Document,p:Path,a):
 raw=read(p);data=None
 try:
  if p.suffix.lower()=='.json':data=json.loads(raw)
  else:
   import yaml;data=yaml.safe_load(raw)
 except Exception as e:callout(doc,'ESTRUTURA NÃO INTERPRETADA',repr(e),AMBER)
 if data is not None:table(doc,['Chave','Valor'],[(k,v[:1000]) for k,v in flatten(data,limit=a.max_struct)],font=6.5)
 doc.add_paragraph('Conteúdo textual:',style='Heading 3');add_text(doc,raw,a.root_resolved,a.max_text_chars)
def add_image(doc:Document,p:Path,a,tmp:Path):
 if Image is None:callout(doc,'PILLOW INDISPONÍVEL',p.name);return
 try:
  with Image.open(p) as im:
   w,h=im.size;fmt=im.format;mode=im.mode;frames=getattr(im,'n_frames',1);x=im.convert('RGB');x.thumbnail((a.image_pixels,a.image_pixels));q=tmp/(hashlib.sha1(str(p).encode()).hexdigest()[:12]+'.jpg');x.save(q,quality=88,optimize=True)
  table(doc,['Campo','Valor'],[('Dimensões',f'{w} × {h}'),('Formato/modo',f'{fmt}/{mode}'),('Frames',frames)])
  par=doc.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER;par.add_run().add_picture(str(q),width=Inches(a.image_inches));cap=doc.add_paragraph(p.relative_to(a.root_resolved).as_posix());cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.runs[0].italic=True
 except Exception as e:callout(doc,'FALHA AO INCORPORAR IMAGEM',repr(e),RED)
def add_notebook(doc:Document,p:Path,a):
 try:nb=json.loads(p.read_text(encoding='utf-8-sig'))
 except Exception as e:callout(doc,'NOTEBOOK INVÁLIDO',repr(e),RED);return
 n=0
 for i,c in enumerate(nb.get('cells',[]),1):
  if c.get('cell_type')!='code' or not c.get('outputs'):continue
  n+=1;doc.add_paragraph(f'Célula {i} · execution_count={c.get("execution_count")}',style='Heading 3');texts=[]
  for o in c['outputs']:
   k=o.get('output_type')
   if k=='stream':texts.append(''.join(o.get('text',[])))
   elif k=='error':texts.append(f'ERRO {o.get("ename")}: {o.get("evalue")}\n'+'\n'.join(o.get('traceback',[])))
   elif k in {'execute_result','display_data'}:
    d=o.get('data',{});v=d.get('text/plain');texts.extend([''.join(v) if isinstance(v,list) else str(v)] if v else []);texts.extend(f'[saída não textual: {m}]' for m in d if m!='text/plain')
  add_text(doc,'\n\n'.join(texts),a.root_resolved,a.max_cell_chars)
 callout(doc,'RESUMO DO NOTEBOOK',f'Células totais: {len(nb.get("cells",[]))}; células com outputs: {n}.')
def add_data(doc:Document,p:Path,a):
 rows=[('Formato',p.suffix.lower()),('Tamanho',size(p.stat().st_size)),('SHA-256',sha(p))]
 if pd is not None:
  try:df=pd.read_parquet(p).head(a.parquet_rows);rows.extend([('Colunas',len(df.columns)),('Prévia',len(df))]);table(doc,['Campo','Valor'],rows);table(doc,[str(x) for x in df.columns[:a.cols_docx]],[tuple('<NA>' if pd.isna(v) else str(v)[:300] for v in r) for _,r in df.iloc[:a.parquet_rows,:a.cols_docx].iterrows()],font=6.5);return
  except Exception as e:rows.append(('Leitura',repr(e)))
 table(doc,['Campo','Valor'],rows)
def setup(doc:Document):
 s=doc.sections[0];s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.65);s.right_margin=Inches(.65);doc.styles['Normal'].font.name='Aptos';doc.styles['Normal'].font.size=Pt(9)
 for n,z,c in [('Title',24,NAVY),('Heading 1',16,NAVY),('Heading 2',13,BLUE),('Heading 3',10,BLUE)]:st=doc.styles[n];st.font.name='Aptos Display';st.font.size=Pt(z);st.font.color.rgb=RGBColor.from_string(c);st.font.bold=True
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.add_run('Compilação local de resultados - tools/read_outputs.py').font.size=Pt(7)
def args():
 p=argparse.ArgumentParser(description='Gera um único DOCX autocontido com resultados do projeto.');p.add_argument('--root',type=Path,default=Path.cwd());p.add_argument('--output-dir',type=Path,default=Path('.local-analysis/compiled_results'));p.add_argument('--name',default='RESULTADOS_COMPLETOS');p.add_argument('--no-notebook-outputs',dest='notebooks',action='store_false');p.set_defaults(notebooks=True);p.add_argument('--max-text-chars',type=int,default=300000);p.add_argument('--max-cell-chars',type=int,default=40000);p.add_argument('--max-rows-read',type=int,default=500000);p.add_argument('--rows-docx',type=int,default=200);p.add_argument('--cols-docx',type=int,default=30);p.add_argument('--summary-cols',type=int,default=300);p.add_argument('--top-values',type=int,default=15);p.add_argument('--max-struct',type=int,default=12000);p.add_argument('--image-inches',type=float,default=6.7);p.add_argument('--image-pixels',type=int,default=2200);p.add_argument('--parquet-rows',type=int,default=30);return p.parse_args()
def main():
 a=args();root=repo_root(a.root);a.root_resolved=root;out=(a.output_dir if a.output_dir.is_absolute() else root/a.output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
 for x in out.iterdir():
  if x.is_file():x.unlink()
  elif x.is_dir():import shutil;shutil.rmtree(x)
 files=scan(root,a.notebooks);doc=Document();setup(doc);p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Resultados completos do projeto');q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run('Arquivo único para análise por IA: textos, tabelas, imagens, manifestos, modelos catalogados e outputs de notebooks').italic=True
 meta=[('Gerado em',datetime.now(timezone.utc).replace(microsecond=0).isoformat()),('Repositório',root.name),('Branch',git(root,'branch','--show-current')),('Commit',git(root,'rev-parse','HEAD')),('Árvore Git','limpa' if not git(root,'status','--porcelain') else 'com alterações locais'),('Python',sys.version.split()[0]),('Plataforma',platform.platform())];table(doc,['Campo','Valor'],meta);callout(doc,'COMO USAR','Envie somente este DOCX à IA. Todos os materiais encontrados estão incorporados. A compilação não recalcula resultados e pode reunir evidências provisórias; sempre cruzar caminho, SHA-256, commit, manifesto e configuração.',AMBER)
 c=Counter(kind(x) for x in files);doc.add_heading('1. Inventário',1);doc.add_paragraph(f'Arquivos encontrados: {len(files)}. '+'; '.join(f'{k}: {v}' for k,v in sorted(c.items()))+'.');table(doc,['#','Caminho','Tipo','Tamanho','SHA-256'],[(i,x.relative_to(root).as_posix(),kind(x),size(x.stat().st_size),sha(x)[:16]+'…') for i,x in enumerate(files,1)],font=6.5);doc.add_page_break();doc.add_heading('2. Conteúdo compilado',1);errors=[]
 with tempfile.TemporaryDirectory(prefix='read_outputs_') as td:
  for i,x in enumerate(files,1):
   rel=x.relative_to(root).as_posix();k=kind(x);print(f'[read_outputs] {i}/{len(files)} {rel}');doc.add_heading(f'2.{i} {rel}',2);table(doc,['Campo','Valor'],[('Tipo',k),('Extensão',x.suffix.lower()),('Tamanho',size(x.stat().st_size)),('Modificado UTC',datetime.fromtimestamp(x.stat().st_mtime,timezone.utc).isoformat()),('SHA-256',sha(x))])
   try:
    if k=='tabela':add_csv(doc,x,a)
    elif k=='texto':add_text(doc,read(x),root,a.max_text_chars)
    elif k=='estruturado':add_struct(doc,x,a)
    elif k=='imagem':add_image(doc,x,a,Path(td))
    elif k=='notebook':add_notebook(doc,x,a)
    elif k=='dados':add_data(doc,x,a)
    elif k=='modelo':callout(doc,'ARTEFATO DE MODELO',f'Binário não desserializado. Tamanho: {size(x.stat().st_size)}. SHA-256: {sha(x)}.')
   except Exception as e:errors.append((rel,repr(e)));callout(doc,'ERRO AO COMPILAR',repr(e),RED)
  if errors:doc.add_heading('3. Erros',1);[doc.add_paragraph(f'{r}: {e}',style='List Bullet') for r,e in errors]
  target=out/(a.name+'.docx');doc.save(target)
 print(f'[read_outputs] Arquivo único: {target}');print(f'[read_outputs] Arquivos: {len(files)}; erros: {len(errors)}');return 1 if errors else 0
if __name__=='__main__':raise SystemExit(main())
